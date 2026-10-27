"""Arma el Excel de Benford: datos, conteo con gráfico y resultado."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill

from benford.ley import BENFORD, chi_cuadrado, clasificar, contar, primer_digito
from config import CHI_BOT, CHI_POSIBLE_BOT, escritorio

COLORES = {"Humano": "C6EFCE", "Posible bot": "FFEB9C", "Bot": "FFC7CE"}


def _encabezado(hoja, columnas: list[str]):
    """Escribe la fila de títulos en negrita con fondo azul y ajusta el ancho."""
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2F5D8F")
        hoja.column_dimensions[celda.column_letter].width = max(14, len(str(celda.value)) + 4)


def generar_excel(datos: dict) -> str:
    """Crea el Excel en el Escritorio y devuelve el resultado: Humano, Posible bot o Bot."""
    usuario = datos["cuenta"]
    libro = Workbook()

    # Hoja 1: las cuentas con sus números y su primer dígito
    hoja = libro.active
    hoja.title = "Datos"
    _encabezado(hoja, ["Cuenta analizada", "Lista", "Usuario", "Seguidores", "Seguidos",
                       "1er dígito seguidores", "1er dígito seguidos"])
    numeros = []
    for c in datos["cuentas"]:
        hoja.append([usuario, datos["lista"], c["usuario"], c["seguidores"], c["seguidos"],
                     primer_digito(c["seguidores"]), primer_digito(c["seguidos"])])
        numeros += [c["seguidores"], c["seguidos"]]

    # Hoja 2: cuántos 1, 2, ... 9, comparado con Benford, y el gráfico
    conteo = contar(numeros)
    total = sum(conteo.values())
    hoja = libro.create_sheet("Conteo")
    _encabezado(hoja, ["Dígito", "Conteo", "% real", "% Benford", "Diferencia", "Tendencia"])
    for d in range(1, 10):
        real = conteo[d] / total * 100 if total else 0
        esperado = BENFORD[d] * 100
        hoja.append([d, conteo[d], round(real, 1), round(esperado, 1),
                     round(real - esperado, 1), "↑ arriba" if real > esperado else "↓ abajo"])

    digitos = Reference(hoja, min_col=1, min_row=2, max_row=10)
    barras = BarChart()                                     # barras = Benford (la base)
    barras.title = "Ley de Benford vs. datos reales"
    barras.x_axis.title = "Primer dígito"
    barras.y_axis.title = "%"
    barras.add_data(Reference(hoja, min_col=4, min_row=1, max_row=10), titles_from_data=True)
    barras.set_categories(digitos)
    linea = LineChart()                                     # línea = lo que contamos
    linea.add_data(Reference(hoja, min_col=3, min_row=1, max_row=10), titles_from_data=True)
    linea.set_categories(digitos)
    barras += linea                                         # la línea va encima de las barras
    barras.x_axis.delete = False
    barras.y_axis.delete = False
    barras.width, barras.height = 18, 9
    hoja.add_chart(barras, "H2")

    # Hoja 3: el chi-cuadrado y la conclusión
    chi = chi_cuadrado(conteo)
    resultado = clasificar(chi)
    hoja = libro.create_sheet("Resultado")
    _encabezado(hoja, ["Dato", "Valor"])
    for fila in [["Cuenta", f"@{usuario}"],
                 ["Lista analizada", datos["lista"]],
                 ["Cuentas revisadas", len(datos["cuentas"])],
                 ["Números usados", total],
                 ["Chi-cuadrado", round(chi, 2)],
                 ["Límite posible bot (95 %)", CHI_POSIBLE_BOT],
                 ["Límite bot (99 %)", CHI_BOT],
                 ["Resultado", resultado]]:
        hoja.append(fila)
    celda = hoja.cell(row=hoja.max_row, column=2)
    celda.font = Font(bold=True)
    celda.fill = PatternFill("solid", fgColor=COLORES[resultado])

    carpeta = escritorio() / f"Benford_{usuario}"
    carpeta.mkdir(exist_ok=True)
    archivo = carpeta / f"benford_{usuario}.xlsx"
    libro.save(archivo)
    print(f"Chi-cuadrado: {chi:.2f} -> {resultado}")
    print(f"Excel guardado en: {archivo}")
    return resultado